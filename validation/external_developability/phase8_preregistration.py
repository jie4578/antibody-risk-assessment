"""Phase 8A: freeze a label-free GDPa3 sequence population before inference.

Only HIC-cell availability is used to apply the preregistered complete-case
eligibility rule. HIC measurement magnitudes are never returned, serialized,
or used to choose rows. This module does not import or invoke an ML runtime.
"""

from __future__ import annotations

import hashlib
import posixpath
import re
import zipfile
from pathlib import Path
from typing import Collection
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook

from validation.schemas.dataset_schema import assess_sequences, sequence_hash


EXPECTED_RAW_SHA256 = "06daa55cb609278574d008c5509d90a63f9a8f1aa3949f34fc50a2e654b008f7"
EXPECTED_EVALUATION_N = 79
AVERAGE_SHEET = "Assay Data - average"
SEQUENCE_SHEET = "Sequences"
HIC_FIELD = "hic_rt"
HIC_AVERAGE_FIELD = "hic_rt_avg"
HIC_UNIT = "minutes"
SEQUENCE_INPUT_COLUMNS = ("antibody_id", "vh_protein_sequence", "lc_protein_sequence")
SNAPSHOT_COLUMNS = ("antibody_id", "VH", "VL", "VH_hash", "VL_hash", "paired_hash")

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _column_from_cell_ref(cell_ref: str) -> str:
    match = re.fullmatch(r"([A-Z]+)[0-9]+", cell_ref)
    if not match:
        raise ValueError(f"Invalid XLSX cell reference: {cell_ref!r}")
    return match.group(1)


def _worksheet_part(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    rel_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rel_root.findall(f"{{{_PKG_REL_NS}}}Relationship")
    }
    for sheet in workbook_root.findall(f".//{{{_MAIN_NS}}}sheet"):
        if sheet.attrib.get("name") == sheet_name:
            rel_id = sheet.attrib.get(f"{{{_REL_NS}}}id")
            if rel_id not in rel_targets:
                break
            target = rel_targets[rel_id].lstrip("/")
            if not target.startswith("xl/"):
                target = posixpath.normpath(posixpath.join("xl", target))
            return target
    raise ValueError(f"XLSX worksheet not found: {sheet_name}")


def nonblank_cell_row_numbers(
    workbook_path: str | Path,
    *,
    sheet_name: str,
    column_letter: str,
) -> set[int]:
    """Return row numbers with a nonblank cell, without interpreting cell data."""

    column_letter = column_letter.upper()
    if not re.fullmatch(r"[A-Z]+", column_letter):
        raise ValueError("column_letter must be an XLSX column name")
    with zipfile.ZipFile(workbook_path) as archive:
        part = _worksheet_part(archive, sheet_name)
        rows_with_value: set[int] = set()
        for _event, element in ET.iterparse(archive.open(part), events=("end",)):
            if element.tag != f"{{{_MAIN_NS}}}row":
                continue
            row_number = int(element.attrib["r"])
            for cell in element.findall(f"{{{_MAIN_NS}}}c"):
                if _column_from_cell_ref(cell.attrib.get("r", "")) != column_letter:
                    continue
                value_node = cell.find(f"{{{_MAIN_NS}}}v")
                # The text is tested only for blankness. It is never parsed as a
                # number, compared, returned, or stored.
                if value_node is not None and value_node.text and value_node.text.strip():
                    rows_with_value.add(row_number)
            element.clear()
    return rows_with_value


def load_gdpa3_phase8_inputs(
    workbook_path: str | Path,
    *,
    expected_n: int = EXPECTED_EVALUATION_N,
) -> tuple[pd.DataFrame, set[str], dict[str, object]]:
    """Read sequence fields and the preregistered HIC availability indicator only."""

    workbook_path = Path(workbook_path)
    if sha256_file(workbook_path) != EXPECTED_RAW_SHA256:
        raise ValueError("GDPa3 raw workbook SHA256 differs from the audited source")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if SEQUENCE_SHEET not in workbook.sheetnames or AVERAGE_SHEET not in workbook.sheetnames:
            raise ValueError("GDPa3 workbook is missing the preregistered sheets")

        sequence_ws = workbook[SEQUENCE_SHEET]
        sequence_header = ["" if value is None else str(value) for value in next(sequence_ws.iter_rows(max_row=1, values_only=True))]
        if not set(SEQUENCE_INPUT_COLUMNS).issubset(sequence_header):
            raise ValueError("GDPa3 Sequences sheet is missing required sequence columns")

        average_ws = workbook[AVERAGE_SHEET]
        average_header = ["" if value is None else str(value) for value in next(average_ws.iter_rows(max_row=1, values_only=True))]
        if "antibody_id" not in average_header or HIC_AVERAGE_FIELD not in average_header:
            raise ValueError("GDPa3 average sheet does not contain the frozen ID/HIC fields")
        id_index = average_header.index("antibody_id")
        hic_index = average_header.index(HIC_AVERAGE_FIELD)

        definitions_ws = workbook["Definitions"] if "Definitions" in workbook.sheetnames else None
        if definitions_ws is None:
            raise ValueError("GDPa3 workbook is missing the Definitions sheet")
        definition_rows = list(definitions_ws.iter_rows(values_only=True))
        hic_definitions = [
            str(row[1])
            for row in definition_rows
            if len(row) > 1 and str(row[0]) == HIC_FIELD and row[1] is not None
        ]
        if len(hic_definitions) != 1 or "retention time" not in hic_definitions[0].lower() or "minute" not in hic_definitions[0].lower():
            raise ValueError("GDPa3 Definitions do not unambiguously define hic_rt as retention time in minutes")

        tidy_header = [
            "" if value is None else str(value)
            for value in next(workbook["Assay Data - tidy format"].iter_rows(max_row=1, values_only=True))
        ] if "Assay Data - tidy format" in workbook.sheetnames else []

        # Workbook rows and IDs are read independently of the HIC cell values.
        average_ids_by_row: dict[int, str] = {}
        for row_number, row in enumerate(
            average_ws.iter_rows(min_row=2, min_col=id_index + 1, max_col=id_index + 1, values_only=True),
            start=2,
        ):
            if row and row[0] is not None and str(row[0]) != "":
                average_ids_by_row[row_number] = str(row[0])

        hic_column_letter = ""
        n = hic_index + 1
        while n:
            n, remainder = divmod(n - 1, 26)
            hic_column_letter = chr(65 + remainder) + hic_column_letter

        source_rows = list(
            sequence_ws.iter_rows(min_row=2, values_only=True)
        )
        id_index_seq = sequence_header.index("antibody_id")
        vh_index = sequence_header.index("vh_protein_sequence")
        vl_index = sequence_header.index("lc_protein_sequence")
        sequence_records = []
        for row in source_rows:
            if row and row[id_index_seq] is not None:
                sequence_records.append(
                    {
                        "antibody_id": str(row[id_index_seq]),
                        "vh_protein_sequence": str(row[vh_index]),
                        "lc_protein_sequence": str(row[vl_index]),
                    }
                )

        nonblank_hic_rows = nonblank_cell_row_numbers(
            workbook_path,
            sheet_name=AVERAGE_SHEET,
            column_letter=hic_column_letter,
        )
        eligible_ids = {
            average_ids_by_row[row_number]
            for row_number in nonblank_hic_rows
            if row_number in average_ids_by_row
        }
        sequence_ids = {row["antibody_id"] for row in sequence_records}
        if len(sequence_ids) != len(sequence_records):
            raise ValueError("GDPa3 Sequences sheet contains duplicate antibody IDs")
        if len(average_ids_by_row) != len(set(average_ids_by_row.values())):
            raise ValueError("GDPa3 average sheet contains duplicate antibody IDs")
        if eligible_ids - sequence_ids:
            raise ValueError("An HIC-available GDPa3 ID has no sequence record")
        if len(eligible_ids) != expected_n:
            raise ValueError(f"Expected {expected_n} HIC-available records; found {len(eligible_ids)}")

        qc_like_fields = [
            field
            for field in sequence_header + average_header + tidy_header
            if any(term in field.lower() for term in ("qc", "quality_control", "pass_fail", "disqual"))
        ]
        if qc_like_fields:
            raise ValueError("GDPa3 contains QC-like fields requiring explicit protocol review")

        return pd.DataFrame(sequence_records, columns=SEQUENCE_INPUT_COLUMNS), eligible_ids, {
            "source_sheet": AVERAGE_SHEET,
            "source_column": HIC_AVERAGE_FIELD,
            "underlying_endpoint": HIC_FIELD,
            "unit": HIC_UNIT,
            "aggregate_definition": "Use the workbook-supplied *_avg field as-is; do not recompute. The workbook does not document a more detailed aggregation formula.",
            "qc_rule": "No workbook-defined disqualifying QC state is present in the audited sequence or average-sheet fields; do not infer QC exclusions from assay magnitudes or replicate summaries.",
            "sequence_rows": len(sequence_records),
            "hic_available_n": len(eligible_ids),
            "hic_missing_n": len(average_ids_by_row) - len(eligible_ids),
            "hic_magnitude_used_for_selection": False,
            "qc_like_fields": qc_like_fields,
        }
    finally:
        workbook.close()


def build_sequence_only_snapshot(
    sequence_frame: pd.DataFrame,
    eligible_ids: Collection[str],
    *,
    expected_n: int = EXPECTED_EVALUATION_N,
) -> pd.DataFrame:
    """Build the exact inference table from sequence-only source columns."""

    if tuple(sequence_frame.columns) != SEQUENCE_INPUT_COLUMNS:
        raise ValueError("Snapshot builder accepts only antibody_id and the two source sequence columns")
    rows = sequence_frame.to_dict(orient="records")
    ids = [str(row["antibody_id"]) for row in rows]
    if any(not record_id or record_id != record_id.strip() for record_id in ids):
        raise ValueError("GDPa3 antibody IDs must be nonempty and preserved without whitespace changes")
    if len(ids) != len(set(ids)):
        raise ValueError("GDPa3 sequence input contains duplicate antibody IDs")

    eligible = {str(record_id) for record_id in eligible_ids}
    source_by_id = {str(row["antibody_id"]): row for row in rows}
    if eligible - set(source_by_id):
        raise ValueError("An eligible GDPa3 ID is missing from the sequence input")
    if len(eligible) != expected_n:
        raise ValueError(f"Expected {expected_n} eligible records; found {len(eligible)}")

    snapshot_rows: list[dict[str, str]] = []
    for record_id in sorted(eligible):
        source = source_by_id[record_id]
        raw_vh = str(source["vh_protein_sequence"])
        raw_vl = str(source["lc_protein_sequence"])
        vh, vl, status, reason = assess_sequences(raw_vh, raw_vl)
        if status != "VALID":
            raise ValueError(f"Eligible GDPa3 record {record_id!r} has invalid paired sequence: {reason}")
        if vh != raw_vh or vl != raw_vl:
            raise ValueError(f"Eligible GDPa3 record {record_id!r} would require sequence normalization")
        snapshot_rows.append(
            {
                "antibody_id": record_id,
                "VH": vh,
                "VL": vl,
                "VH_hash": hashlib.sha256(vh.encode("ascii")).hexdigest(),
                "VL_hash": hashlib.sha256(vl.encode("ascii")).hexdigest(),
                "paired_hash": sequence_hash(vh, vl),
            }
        )
    snapshot = pd.DataFrame(snapshot_rows, columns=SNAPSHOT_COLUMNS)
    validate_sequence_snapshot_schema(snapshot)
    return snapshot


def validate_sequence_snapshot_schema(snapshot: pd.DataFrame) -> None:
    """Enforce an exact label-free inference snapshot schema and unique keys."""

    if tuple(snapshot.columns) != SNAPSHOT_COLUMNS:
        raise ValueError("GDPa3 inference snapshot has unexpected or label-bearing columns")
    if snapshot["antibody_id"].duplicated().any() or snapshot["paired_hash"].duplicated().any():
        raise ValueError("GDPa3 inference snapshot IDs and paired hashes must be unique")


def write_sequence_snapshot(snapshot: pd.DataFrame, path: str | Path) -> str:
    validate_sequence_snapshot_schema(snapshot)
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    snapshot.to_csv(destination, index=False, encoding="utf-8", lineterminator="\n")
    return sha256_file(destination)

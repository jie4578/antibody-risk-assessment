from __future__ import annotations

import hashlib

import pandas as pd
from openpyxl import Workbook

from validation.external_developability.gdpa3_audit import (
    _sequence_records,
    audit_sequence_sources,
    exact_overlap,
    identity_proxy_clusters,
    inspect_gdpa3_workbook,
    nearest_paired_min_audit,
    sha256_file,
)
from validation.schemas.dataset_schema import sequence_hash


VH_A = "ACDEFGHIKLMNPQRSTVWY"
VL_A = "YWVTSRQPNMLKIHGFEDCA"
VH_B = "CDEFGHIKLMNPQRSTVWYA"
VL_B = "WVTSRQPNMLKIHGFEDCAY"


def _write_small_source(path) -> None:
    wb = Workbook()
    definitions = wb.active
    definitions.title = "Definitions"
    definitions.append(["Column name", "Column description"])
    definitions.append(["vh_protein_sequence", "heavy variable-domain sequence"])
    definitions.append(["vl_protein_sequence", "light variable-domain sequence"])
    definitions.append(["hic_rt", "HIC retention time in minutes"])

    sequences = wb.create_sheet("Sequences")
    sequences.append(
        ["antibody_id", "hc_subtype", "lc_subtype", "vh_protein_sequence", "lc_protein_sequence", "hc_dna_sequence", "lc_dna_sequence"]
    )
    sequences.append(["id-1", "IgG1", "kappa", f" {VH_A.lower()} ", VL_A, None, None])
    sequences.append(["id-2", "IgG1", "kappa", VH_A, VL_A, None, None])
    sequences.append(["id-3", "IgG1", "lambda", VH_B, None, None, None])
    sequences.append(["id-4", "IgG1", "kappa", VH_A + "X", VL_B, None, None])

    tidy = wb.create_sheet("Assay Data - tidy format")
    tidy.append(["antibody_id", "technical_replicate", "hic_rt", "titer"])
    tidy.append(["id-1", 1, 10.0, 100.0])
    tidy.append(["id-1", 2, None, 110.0])
    tidy.append(["id-2", 1, 11.0, None])

    average = wb.create_sheet("Assay Data - average")
    average.append(["antibody_id", "hic_rt_avg", "hic_rt_replicates", "hic_rt_stddev", "titer_avg", "titer_replicates", "titer_stddev"])
    average.append(["id-1", 10.0, 1, None, 105.0, 2, 7.0])
    average.append(["id-2", 11.0, 1, None, None, 0, None])
    average.append(["id-3", None, 0, None, None, 0, None])
    average.append(["id-4", None, 0, None, None, 0, None])

    versioning = wb.create_sheet("Versioning")
    versioning.append(["Date", "Version", "Description"])
    versioning.append(["2026-01-06", "test", "synthetic fixture"])
    wb.save(path)


def test_workbook_structure_sequence_mapping_and_assay_inventory_are_deterministic(tmp_path) -> None:
    path = tmp_path / "synthetic.xlsx"
    _write_small_source(path)
    digest_before = sha256_file(path)

    first, records = inspect_gdpa3_workbook(path)
    second, second_records = inspect_gdpa3_workbook(path)

    assert first == second
    pd.testing.assert_frame_equal(records, second_records)
    assert sha256_file(path) == digest_before
    assert first["workbook"]["sheet_names"] == [
        "Definitions",
        "Sequences",
        "Assay Data - tidy format",
        "Assay Data - average",
        "Versioning",
    ]
    assert first["sequence_fields"]["VH"] == "vh_protein_sequence"
    assert first["sequence_fields"]["VL"] == "lc_protein_sequence"
    assert first["sequence_fields"]["source_header_difference"]["mismatch"] is True
    assert first["sequence_audit"]["rows"] == 4
    assert first["sequence_audit"]["sequence_status_counts"] == {"VALID": 2, "PARTIAL": 1, "INVALID": 1}
    assert first["sequence_audit"]["unique_valid_paired_hashes"] == 1
    assert first["sequence_audit"]["duplicate_valid_paired_rows"] == 1
    assert first["sequence_audit"]["normalized_sequence_rows"] == 1
    hic = next(item for item in first["assay_inventory"] if item["field"] == "hic_rt")
    assert hic["tidy_nonmissing_measurements"] == 2
    assert hic["average_nonmissing_antibodies"] == 2
    assert hic["replicate_count_nonmissing_antibodies"] == 4
    assert hic["standard_deviation_nonmissing_antibodies"] == 0
    assert first["assay_id_audit"]["Assay Data - average"]["sequence_only_ids"] == 0
    crosswalk = {item["competition_endpoint"]: item for item in first["competition_endpoint_crosswalk"]}
    assert crosswalk["HIC"]["workbook_field"] == "hic_rt"
    assert crosswalk["Tm2"]["workbook_field"] == "tm2_nanodsf"
    assert crosswalk["AC-SINS pH 7.4"]["workbook_field"] == "acsins_dLmax_ph7.4"
    assert first["frozen_model_compatibility"]["hic_esm2_v1"]["protocol_equivalence_confirmed"] is False
    assert first["frozen_model_compatibility"]["developability_esm2_v1"]["compatible"] is False


def test_sequence_hash_is_stable_and_duplicate_rows_are_not_dropped() -> None:
    frame = pd.DataFrame(
        [
            {"record_id": "one", "VH": f" {VH_A.lower()} ", "VL": VL_A},
            {"record_id": "two", "VH": VH_A, "VL": VL_A},
        ]
    )
    rows = _sequence_records(frame, id_column="record_id")
    assert len(rows) == 2
    assert rows["sequence_hash"].nunique() == 1
    assert rows.iloc[0]["sequence_hash"] == sequence_hash(VH_A, VL_A)
    assert rows.iloc[0]["sequence_hash"] == rows.iloc[1]["sequence_hash"]
    assert hashlib.sha256(VH_A.encode("ascii")).hexdigest() == rows.iloc[0]["vh_hash"]


def test_exact_overlap_counts_individual_chains_and_paired_sequences_separately() -> None:
    query = _sequence_records(pd.DataFrame([{"VH": VH_A, "VL": VL_A}]))
    comparison = _sequence_records(
        pd.DataFrame([{"VH": VH_A, "VL": VL_B}, {"VH": VH_B, "VL": VL_A}])
    )
    assert exact_overlap(query, comparison) == {"VH": 1, "VL": 1, "paired": 0}


def test_nearest_identity_and_novelty_bins_are_reproducible() -> None:
    query = _sequence_records(pd.DataFrame([{"VH": VH_A, "VL": VL_A}, {"VH": VH_B, "VL": VL_B}]))
    reference = _sequence_records(pd.DataFrame([{"VH": VH_A, "VL": VL_A}]))
    first = nearest_paired_min_audit(query, reference)
    second = nearest_paired_min_audit(query, reference)
    assert first == second
    assert first["nearest_paired_min"]["n"] == 2
    assert sum(first["novelty_bins"].values()) == 2
    assert first["nearest_paired_min"]["max"] == 1.0


def test_identity_proxy_clustering_is_repeatable_and_threshold_sensitive() -> None:
    records = _sequence_records(
        pd.DataFrame(
            [
                {"VH": VH_A, "VL": VL_A},
                {"VH": VH_B, "VL": VL_B},
                {"VH": "M" + VH_A[1:], "VL": "M" + VL_A[1:]},
            ]
        )
    )
    loose = identity_proxy_clusters(records, 0.90)
    strict = identity_proxy_clusters(records, 1.0)
    assert loose == identity_proxy_clusters(records, 0.90)
    assert loose["clusters"] < strict["clusters"]
    assert strict["clusters"] == 3
    assert strict["sequence_clusters_are_germline_families"] is False


def test_split_overlap_audit_uses_hash_membership_and_no_outcome_fields() -> None:
    gdpa3 = _sequence_records(pd.DataFrame([{"VH": VH_A, "VL": VL_A}]))
    split = pd.DataFrame(
        [
            {"sequence_hash": sequence_hash(VH_B, VL_B), "split": "TRAIN"},
            {"sequence_hash": sequence_hash(VH_A, VL_A), "split": "TEST"},
        ]
    )
    result = audit_sequence_sources(
        gdpa3,
        pd.DataFrame([{"antibody_id": "jain", "VH": VH_B, "VL": VL_B}]),
        pd.DataFrame(
            [
                {"record_id": "train", "VH": VH_B, "VL": VL_B, "sequence_hash": sequence_hash(VH_B, VL_B)},
                {"record_id": "test", "VH": VH_A, "VL": VL_A, "sequence_hash": sequence_hash(VH_A, VL_A)},
            ]
        ),
        split,
        pd.DataFrame([{"VH": VH_B, "VL": VL_B}]),
    )
    assert result["exact_overlap"]["AIntibody_frozen_splits_paired_hashes_only"]["TEST"] == 1
    assert result["exact_overlap"]["AIntibody_frozen_splits_paired_hashes_only"]["TRAIN"] == 0
    assert result["test_label_access"].startswith("NO")
    assert result["source_hash_integrity"]["AIntibody_source_sequence_hash_mismatches"] == 0


def test_phase7c1_runner_is_sequence_only_and_never_opens_test_labels() -> None:
    from pathlib import Path

    runner = Path(__file__).resolve().parents[2] / "run_gdpa3_phase7c1_audit.py"
    source = runner.read_text(encoding="utf-8")
    assert "sealed_test_labels" not in source
    assert "usecols=[\"sequence_hash\", \"split\"]" in source
    assert "usecols=[\"record_id\", \"VH\", \"VL\", \"sequence_hash\"]" in source
    assert "HIC\"" not in source


def test_raw_workbook_path_ignore_rule_is_narrow() -> None:
    from pathlib import Path

    repository = Path(__file__).resolve().parents[3]
    rules = (repository / ".gitignore").read_text(encoding="utf-8").splitlines()
    assert "validation/external_developability/raw/" in rules
    assert "*.xlsx" not in rules

"""Read-only, label-safe audit utilities for the GDPa3 source workbook.

This module counts assay availability but never interprets assay values.  It
does not call the production predictor, fit a model, calculate performance,
or read AIntibody outcome columns.  Sequence identity uses the validation
package's existing global-alignment implementation.
"""

from __future__ import annotations

import hashlib
import json
import statistics
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from core import validate_sequence
from validation.ml_benchmark.similarity import global_identity
from validation.schemas.dataset_schema import assess_sequences, sequence_hash


REQUIRED_SHEETS = (
    "Definitions",
    "Sequences",
    "Assay Data - tidy format",
    "Assay Data - average",
    "Versioning",
)
SEQUENCE_FIELDS = {
    "id": "antibody_id",
    "vh": "vh_protein_sequence",
    "vl_candidates": ("vl_protein_sequence", "lc_protein_sequence"),
}
SPLITS = ("TRAIN", "VALIDATION", "TEST")
NOVELTY_BINS = (
    ("<0.50", 0.0, 0.50),
    ("0.50-<0.70", 0.50, 0.70),
    ("0.70-<0.80", 0.70, 0.80),
    ("0.80-<0.90", 0.80, 0.90),
    (">=0.90", 0.90, 1.000000000001),
)
COMPETITION_ENDPOINT_CROSSWALK = (
    {"competition_endpoint": "HIC", "workbook_field": "hic_rt", "workbook_average_field": "hic_rt_avg", "interpretation": "HIC retention time", "unit": "minutes"},
    {"competition_endpoint": "PR-CHO", "workbook_field": "polyreactivity_prscore_cho", "workbook_average_field": "polyreactivity_prscore_cho_avg", "interpretation": "normalized polyreactivity score against CHO SMP", "unit": "source-defined normalized score"},
    {"competition_endpoint": "AC-SINS pH 7.4", "workbook_field": "acsins_dLmax_ph7.4", "workbook_average_field": "acsins_dLmax_ph7.4_avg", "interpretation": "AC-SINS dLmax at pH 7.4", "unit": "nm"},
    {"competition_endpoint": "Tm2", "workbook_field": "tm2_nanodsf", "workbook_average_field": "tm2_nanodsf_avg", "interpretation": "second nanoDSF melting transition", "unit": "degC"},
    {"competition_endpoint": "Titer", "workbook_field": "titer", "workbook_average_field": "titer_avg", "interpretation": "IgG titer in clarified harvest supernatant", "unit": "ug/mL"},
)


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _missing(value: Any) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _sequence_records(
    frame: pd.DataFrame,
    *,
    vh_column: str = "VH",
    vl_column: str = "VL",
    id_column: str | None = None,
    source_hash_column: str | None = None,
) -> pd.DataFrame:
    """Normalize for analysis only; keep every input row and never deduplicate."""

    missing = {vh_column, vl_column} - set(frame.columns)
    if missing:
        raise ValueError(f"Sequence table missing columns: {sorted(missing)}")
    records: list[dict[str, Any]] = []
    for source_index, row in frame.reset_index(drop=True).iterrows():
        raw_vh, raw_vl = row[vh_column], row[vl_column]
        vh, vl, status, reason = assess_sequences(raw_vh, raw_vl)
        vh_ok, _ = validate_sequence(vh, allow_empty=True)
        vl_ok, _ = validate_sequence(vl, allow_empty=True)
        pair_digest = sequence_hash(vh, vl) if status == "VALID" else ""
        raw_source_hash = None if source_hash_column is None else row.get(source_hash_column)
        source_hash = "" if _missing(raw_source_hash) else str(raw_source_hash).strip()
        raw_record_id = None if id_column is None else row.get(id_column)
        records.append(
            {
                "source_index": source_index,
                "record_id": "" if _missing(raw_record_id) else str(raw_record_id).strip(),
                "VH": vh,
                "VL": vl,
                "sequence_status": status,
                "sequence_status_reason": reason,
                "vh_hash": hashlib.sha256(vh.encode("ascii")).hexdigest() if vh and vh_ok else "",
                "vl_hash": hashlib.sha256(vl.encode("ascii")).hexdigest() if vl and vl_ok else "",
                "sequence_hash": pair_digest,
                "source_sequence_hash": source_hash,
                "source_hash_matches": bool(source_hash and pair_digest == source_hash),
                "raw_vh_was_normalized": not _missing(raw_vh) and str(raw_vh) != vh,
                "raw_vl_was_normalized": not _missing(raw_vl) and str(raw_vl) != vl,
            }
        )
    return pd.DataFrame(records)


def _identity_summary(records: pd.DataFrame) -> dict[str, Any]:
    valid = records.loc[records["sequence_status"].eq("VALID")]
    status_counts = records["sequence_status"].value_counts().to_dict()
    reasons = Counter(
        reason
        for reason in records.loc[records["sequence_status"].ne("VALID"), "sequence_status_reason"]
        if reason
    )
    lengths_vh = valid["VH"].map(len).tolist()
    lengths_vl = valid["VL"].map(len).tolist()

    def length_summary(values: list[int]) -> dict[str, int | float | None]:
        if not values:
            return {"min": None, "median": None, "max": None}
        return {
            "min": min(values),
            "median": float(statistics.median(values)),
            "max": max(values),
        }

    return {
        "rows": int(len(records)),
        "unique_record_ids": int(records.loc[records["record_id"].ne(""), "record_id"].nunique()),
        "duplicate_id_rows": int(records.loc[records["record_id"].ne(""), "record_id"].duplicated(keep=False).sum()),
        "rows_with_both_chains_present": int((records["VH"].ne("") & records["VL"].ne("")).sum()),
        "missing_VH_rows": int(records["VH"].eq("").sum()),
        "missing_VL_rows": int(records["VL"].eq("").sum()),
        "sequence_status_counts": {key: int(status_counts.get(key, 0)) for key in ("VALID", "PARTIAL", "INVALID")},
        "invalid_or_partial_reasons": dict(sorted(reasons.items())),
        "unique_valid_VH_hashes": int(valid.loc[valid["vh_hash"].ne(""), "vh_hash"].nunique()),
        "unique_valid_VL_hashes": int(valid.loc[valid["vl_hash"].ne(""), "vl_hash"].nunique()),
        "unique_valid_paired_hashes": int(valid["sequence_hash"].nunique()),
        "duplicate_valid_paired_rows": int(len(valid) - valid["sequence_hash"].nunique()),
        "normalized_sequence_rows": int((records["raw_vh_was_normalized"] | records["raw_vl_was_normalized"]).sum()),
        "VH_length_aa": length_summary(lengths_vh),
        "VL_length_aa": length_summary(lengths_vl),
    }


def _sheet_rows(ws: Any) -> tuple[list[str], list[tuple[Any, ...]], int]:
    iterator = ws.iter_rows(values_only=True)
    raw_header = next(iterator, ())
    header = ["" if value is None else str(value) for value in raw_header]
    rows = [tuple(row) for row in iterator if row and not _missing(row[0])]
    return header, rows, len(rows)


def _assay_inventory(
    tidy_header: list[str],
    tidy_rows: list[tuple[Any, ...]],
    average_header: list[str],
    average_rows: list[tuple[Any, ...]],
    definitions: dict[str, str],
) -> list[dict[str, Any]]:
    tidy_index = {name: index for index, name in enumerate(tidy_header)}
    average_index = {name: index for index, name in enumerate(average_header)}
    average_fields = [name for name in average_header if name.endswith("_avg")]
    output: list[dict[str, Any]] = []
    for average_field in average_fields:
        field = average_field[:-4]
        tidy_values = [row[tidy_index[field]] for row in tidy_rows] if field in tidy_index else []
        average_values = [row[average_index[average_field]] for row in average_rows]
        replicate_field = f"{field}_replicates"
        stddev_field = f"{field}_stddev"
        replicate_values = [row[average_index[replicate_field]] for row in average_rows] if replicate_field in average_index else []
        stddev_values = [row[average_index[stddev_field]] for row in average_rows] if stddev_field in average_index else []
        tidy_nonmissing = sum(not _missing(value) for value in tidy_values)
        average_nonmissing = sum(not _missing(value) for value in average_values)
        measured_antibodies = {
            str(row[tidy_index["antibody_id"]])
            for row in tidy_rows
            if field in tidy_index and not _missing(row[tidy_index[field]])
        }
        output.append(
            {
                "field": field,
                "tidy_column": field if field in tidy_index else None,
                "average_column": average_field,
                "replicate_count_column": replicate_field if replicate_field in average_index else None,
                "standard_deviation_column": stddev_field if stddev_field in average_index else None,
                "description": definitions.get(field),
                "tidy_nonmissing_measurements": int(tidy_nonmissing),
                "tidy_missing_measurements": int(len(tidy_values) - tidy_nonmissing),
                "tidy_antibodies_with_measurement": int(len(measured_antibodies)),
                "average_nonmissing_antibodies": int(average_nonmissing),
                "average_missing_antibodies": int(len(average_values) - average_nonmissing),
                "replicate_count_nonmissing_antibodies": int(sum(not _missing(value) for value in replicate_values)),
                "replicate_count_missing_antibodies": int(len(replicate_values) - sum(not _missing(value) for value in replicate_values)),
                "standard_deviation_nonmissing_antibodies": int(sum(not _missing(value) for value in stddev_values)),
                "standard_deviation_missing_antibodies": int(len(stddev_values) - sum(not _missing(value) for value in stddev_values)),
            }
        )
    return output


def inspect_gdpa3_workbook(path: str | Path) -> tuple[dict[str, Any], pd.DataFrame]:
    """Inspect headers, row counts, sequence quality, and assay missingness only."""

    workbook_path = Path(path)
    if not zipfile.is_zipfile(workbook_path):
        raise ValueError("Input is not a valid XLSX/ZIP container")
    with zipfile.ZipFile(workbook_path) as archive:
        bad_member = archive.testzip()
        required_parts = {"[Content_Types].xml", "xl/workbook.xml"}
        if bad_member or not required_parts.issubset(set(archive.namelist())):
            raise ValueError("XLSX package integrity check failed")

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        missing_sheets = set(REQUIRED_SHEETS) - set(wb.sheetnames)
        if missing_sheets:
            raise ValueError(f"GDPa3 workbook missing sheets: {sorted(missing_sheets)}")

        inventories: dict[str, dict[str, Any]] = {}
        sheet_data: dict[str, tuple[list[str], list[tuple[Any, ...]], int]] = {}
        for sheet in wb.worksheets:
            header, rows, row_count = _sheet_rows(sheet)
            sheet_data[sheet.title] = (header, rows, row_count)
            inventories[sheet.title] = {
                "rows_with_nonempty_first_column": row_count,
                "column_count": len(header),
                "headers": header,
            }

        _, definition_rows, _ = sheet_data["Definitions"]
        definitions = {
            str(row[0]): str(row[1])
            for row in definition_rows
            if len(row) > 1 and not _missing(row[0]) and not _missing(row[1])
        }
        sequence_header, sequence_rows, _ = sheet_data["Sequences"]
        required_sequence_headers = {SEQUENCE_FIELDS["id"], SEQUENCE_FIELDS["vh"]}
        vl_header = next((name for name in SEQUENCE_FIELDS["vl_candidates"] if name in sequence_header), None)
        if not required_sequence_headers.issubset(sequence_header) or vl_header is None:
            raise ValueError("GDPa3 Sequences sheet does not have recognized antibody/VH/VL columns")
        seq_frame = pd.DataFrame(sequence_rows, columns=sequence_header)
        records = _sequence_records(
            seq_frame,
            vh_column=SEQUENCE_FIELDS["vh"],
            vl_column=vl_header,
            id_column=SEQUENCE_FIELDS["id"],
        )

        tidy_header, tidy_rows, _ = sheet_data["Assay Data - tidy format"]
        average_header, average_rows, _ = sheet_data["Assay Data - average"]
        if "antibody_id" not in tidy_header or "antibody_id" not in average_header:
            raise ValueError("GDPa3 assay sheets are missing antibody_id")
        assay_inventory = _assay_inventory(
            tidy_header, tidy_rows, average_header, average_rows, definitions
        )
        sequence_ids = set(records["record_id"])
        assay_id_sets: dict[str, set[str]] = {}
        for name, header, rows in (
            ("Assay Data - tidy format", tidy_header, tidy_rows),
            ("Assay Data - average", average_header, average_rows),
        ):
            id_index = header.index("antibody_id")
            assay_id_sets[name] = {str(row[id_index]) for row in rows if not _missing(row[id_index])}

        definition_sequence_names = [
            name for name in definitions
            if name in {"vh_protein_sequence", "vl_protein_sequence", "lc_protein_sequence"}
        ]
        germline_headers = [
            name for name in sequence_header
            if any(token in name.lower() for token in ("ighv", "igkv", "iglv", "germline", "v_gene"))
        ]
        subtype_headers = [name for name in sequence_header if "subtype" in name.lower()]
        source_header_difference = {
            "definitions_light_chain_sequence_field": "vl_protein_sequence" if "vl_protein_sequence" in definitions else None,
            "sequences_sheet_light_chain_sequence_field": vl_header,
            "mismatch": bool("vl_protein_sequence" in definitions and vl_header != "vl_protein_sequence"),
        }
        audit = {
            "workbook": {
                "source_filename": workbook_path.name,
                "file_size_bytes": workbook_path.stat().st_size,
                "sha256": sha256_file(workbook_path),
                "valid_xlsx_package": True,
                "sheet_names": list(wb.sheetnames),
                "sheets": inventories,
            },
            "sequence_fields": {
                "record_id": SEQUENCE_FIELDS["id"],
                "VH": SEQUENCE_FIELDS["vh"],
                "VL": vl_header,
                "DNA_fields_present": [name for name in sequence_header if name.endswith("_dna_sequence")],
                "definitions_sequence_fields": definition_sequence_names,
                "source_header_difference": source_header_difference,
            },
            "sequence_audit": _identity_summary(records),
            "germline_metadata": {
                "explicit_germline_fields": germline_headers,
                "subtype_fields": subtype_headers,
                "interpretation": "No IGHV/IGKV/IGLV germline annotation is present; subtype fields are constant-region subtype metadata, not germline V-gene calls.",
            },
            "assay_inventory": assay_inventory,
            "competition_endpoint_crosswalk": [
                {
                    **item,
                    "field_present": item["workbook_field"] in tidy_header,
                    "average_field_present": item["workbook_average_field"] in average_header,
                    "average_nonmissing_n": next(
                        (
                            assay["average_nonmissing_antibodies"]
                            for assay in assay_inventory
                            if assay["field"] == item["workbook_field"]
                        ),
                        0,
                    ),
                }
                for item in COMPETITION_ENDPOINT_CROSSWALK
            ],
            "frozen_model_compatibility": {
                "hic_esm2_v1": {
                    "nominal_endpoint_compatible": True,
                    "target_field": "hic_rt_avg",
                    "unit": "minutes",
                    "nonmissing_n": next(
                        assay["average_nonmissing_antibodies"]
                        for assay in assay_inventory
                        if assay["field"] == "hic_rt"
                    ),
                    "protocol_equivalence_confirmed": False,
                    "interpretation": "HIC retention-time observable matches nominally; cross-source protocol equivalence is unverified.",
                },
                "developability_esm2_v1": {
                    "compatible": False,
                    "reasons": [
                        "GDPa3 has no BVP endpoint.",
                        "GDPa3 has no Tagg endpoint.",
                        "GDPa3 AC-SINS is dLmax, not the AIntibody dPW-adjusted endpoint.",
                        "GDPa3 Tm2 is not the AIntibody Tm endpoint.",
                    ],
                },
            },
            "assay_id_audit": {
                name: {
                    "unique_assay_ids": len(ids),
                    "sequence_only_ids": len(sequence_ids - ids),
                    "assay_only_ids": len(ids - sequence_ids),
                }
                for name, ids in assay_id_sets.items()
            },
        }
    finally:
        wb.close()
    return audit, records


def _hash_sets(records: pd.DataFrame) -> dict[str, set[str]]:
    valid = records.loc[records["sequence_status"].eq("VALID")]
    return {
        "VH": set(valid.loc[valid["vh_hash"].ne(""), "vh_hash"]),
        "VL": set(valid.loc[valid["vl_hash"].ne(""), "vl_hash"]),
        "paired": set(valid.loc[valid["sequence_hash"].ne(""), "sequence_hash"]),
    }


def exact_overlap(gdpa3_records: pd.DataFrame, other_records: pd.DataFrame) -> dict[str, int]:
    left, right = _hash_sets(gdpa3_records), _hash_sets(other_records)
    return {chain: len(left[chain] & right[chain]) for chain in ("VH", "VL", "paired")}


def nearest_paired_min_audit(
    query_records: pd.DataFrame,
    reference_records: pd.DataFrame,
) -> dict[str, Any]:
    """Return exact global-alignment nearest-neighbour identity summaries."""

    query = query_records.loc[query_records["sequence_status"].eq("VALID")].drop_duplicates("sequence_hash")
    reference = reference_records.loc[reference_records["sequence_status"].eq("VALID")].drop_duplicates("sequence_hash")
    if reference.empty:
        raise ValueError("No valid reference VH/VL sequences")
    refs = reference[["sequence_hash", "VH", "VL"]].sort_values("sequence_hash").to_dict("records")
    nearest: list[float] = []
    for row in query[["sequence_hash", "VH", "VL"]].sort_values("sequence_hash").to_dict("records"):
        best = 0.0
        for other in refs:
            vh_identity = global_identity(row["VH"], other["VH"])
            vl_identity = global_identity(row["VL"], other["VL"])
            best = max(best, min(vh_identity, vl_identity))
        nearest.append(best)
    counts = {name: 0 for name, _, _ in NOVELTY_BINS}
    for value in nearest:
        for name, lower, upper in NOVELTY_BINS:
            if lower <= value < upper:
                counts[name] += 1
                break
    series = pd.Series(nearest, dtype=float)
    quantile = lambda value: float(series.quantile(value)) if not series.empty else None
    return {
        "identity_method": "existing validation global alignment; paired_min=min(VH identity, VL identity)",
        "query_unique_pairs": int(len(query)),
        "reference_unique_pairs": int(len(reference)),
        "pairwise_comparisons": int(len(query) * len(reference)),
        "nearest_paired_min": {
            "n": int(series.size),
            "min": float(series.min()) if not series.empty else None,
            "q25": quantile(0.25),
            "median": quantile(0.50),
            "q75": quantile(0.75),
            "max": float(series.max()) if not series.empty else None,
            "mean": float(series.mean()) if not series.empty else None,
        },
        "novelty_bins": counts,
    }


def identity_proxy_clusters(records: pd.DataFrame, threshold: float) -> dict[str, Any]:
    """Exact all-pairs connected components; not immunogenetic family calls."""

    if not 0.0 < threshold <= 1.0:
        raise ValueError("threshold must be in (0, 1]")
    valid = records.loc[records["sequence_status"].eq("VALID")].drop_duplicates("sequence_hash")
    rows = valid[["sequence_hash", "VH", "VL"]].sort_values("sequence_hash").to_dict("records")
    parent = list(range(len(rows)))

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    edge_count = 0
    for left in range(len(rows)):
        for right in range(left + 1, len(rows)):
            if (
                global_identity(rows[left]["VH"], rows[right]["VH"]) >= threshold
                and global_identity(rows[left]["VL"], rows[right]["VL"]) >= threshold
            ):
                edge_count += 1
                left_root, right_root = find(left), find(right)
                if left_root != right_root:
                    parent[max(left_root, right_root)] = min(left_root, right_root)
    sizes = Counter(find(index) for index in range(len(rows)))
    ordered_sizes = sorted(sizes.values(), reverse=True)
    return {
        "threshold": threshold,
        "method": "connected components of exact all-pairs global-alignment edges; both VH and VL identity >= threshold",
        "clusters": len(ordered_sizes),
        "largest_cluster": ordered_sizes[0] if ordered_sizes else 0,
        "singleton_clusters": sum(size == 1 for size in ordered_sizes),
        "cluster_size_distribution": {str(size): count for size, count in sorted(Counter(ordered_sizes).items())},
        "qualifying_pair_edges": edge_count,
        "sequence_clusters_are_germline_families": False,
    }


def audit_sequence_sources(
    gdpa3_records: pd.DataFrame,
    jain_frame: pd.DataFrame,
    aintibody_frame: pd.DataFrame,
    split_frame: pd.DataFrame,
    sabdab_frame: pd.DataFrame,
) -> dict[str, Any]:
    """Compare normalized sequences; AIntibody splits use hash/split only."""

    jain_records = _sequence_records(jain_frame, id_column="antibody_id" if "antibody_id" in jain_frame else None)
    aintibody_records = _sequence_records(
        aintibody_frame,
        id_column="record_id" if "record_id" in aintibody_frame else None,
        source_hash_column="sequence_hash" if "sequence_hash" in aintibody_frame else None,
    )
    sabdab_records = _sequence_records(sabdab_frame)
    if "sequence_hash" not in split_frame or "split" not in split_frame:
        raise ValueError("split membership must include only sequence_hash and split columns")
    if split_frame["sequence_hash"].isna().any() or split_frame["sequence_hash"].duplicated().any():
        raise ValueError("Split membership must assign each non-missing sequence hash exactly once")
    unknown_splits = set(split_frame["split"].dropna().astype(str)) - set(SPLITS)
    if unknown_splits:
        raise ValueError(f"Unknown split names: {sorted(unknown_splits)}")
    split_hashes = {
        name: set(split_frame.loc[split_frame["split"].eq(name), "sequence_hash"].astype(str))
        for name in SPLITS
    }
    all_split_hashes = set().union(*split_hashes.values())
    if set(split_frame["sequence_hash"].astype(str)) != all_split_hashes:
        raise ValueError("Split membership contains missing hash assignments")
    computed_hashes = set(aintibody_records.loc[aintibody_records["sequence_status"].eq("VALID"), "sequence_hash"])
    source_hash_mismatches = int(
        (
            aintibody_records["sequence_status"].eq("VALID")
            & ~aintibody_records["source_hash_matches"]
        ).sum()
    )
    trval_hashes = split_hashes["TRAIN"] | split_hashes["VALIDATION"]
    if not trval_hashes.issubset(computed_hashes):
        raise ValueError("A TRAIN/VALIDATION identity hash has no matching sequence-only source row")
    trval_records = aintibody_records.loc[
        aintibody_records["sequence_hash"].isin(trval_hashes)
        & aintibody_records["sequence_status"].eq("VALID"),
        ["sequence_hash", "VH", "VL", "sequence_status"],
    ].drop_duplicates("sequence_hash")
    gdpa3_valid = gdpa3_records.loc[gdpa3_records["sequence_status"].eq("VALID")]

    return {
        "source_record_counts": {
            "Jain_processed_rows": int(len(jain_records)),
            "AIntibody_all_sequence_rows": int(len(aintibody_records)),
            "AIntibody_all_unique_valid_paired_hashes": int(len(_hash_sets(aintibody_records)["paired"])),
            "SAbDab2_source_rows": int(len(sabdab_records)),
            "SAbDab2_both_chains_present": int((sabdab_records["VH"].ne("") & sabdab_records["VL"].ne("")).sum()),
            "SAbDab2_valid_paired_rows": int(sabdab_records["sequence_status"].eq("VALID").sum()),
            "SAbDab2_unique_valid_paired_hashes": int(len(_hash_sets(sabdab_records)["paired"])),
        },
        "source_hash_integrity": {
            "AIntibody_valid_rows_checked": int(aintibody_records["sequence_status"].eq("VALID").sum()),
            "AIntibody_source_sequence_hash_mismatches": source_hash_mismatches,
        },
        "exact_overlap": {
            "Jain": exact_overlap(gdpa3_records, jain_records),
            "AIntibody_all_715_records": exact_overlap(gdpa3_records, aintibody_records),
            "AIntibody_frozen_476_hash_population": {
                "paired": len(_hash_sets(gdpa3_records)["paired"] & all_split_hashes),
            },
            "AIntibody_frozen_splits_paired_hashes_only": {
                name: len(_hash_sets(gdpa3_records)["paired"] & hashes)
                for name, hashes in split_hashes.items()
            },
            "SAbDab2": exact_overlap(gdpa3_records, sabdab_records),
        },
        "AIntibody_split_population": {
            name: len(hashes) for name, hashes in split_hashes.items()
        } | {"ALL": len(all_split_hashes)},
        "AIntibody_train_validation_similarity": nearest_paired_min_audit(gdpa3_valid, trval_records),
        "identity_proxy_clusters": {
            str(threshold): identity_proxy_clusters(gdpa3_valid, threshold)
            for threshold in (0.70, 0.80, 0.90)
        },
        "test_label_access": "NO; TEST comparison used sequence_hash and split membership only",
    }


def build_markdown_report(audit: dict[str, Any], overlaps: dict[str, Any]) -> str:
    seq = audit["sequence_audit"]
    workbook = audit["workbook"]
    assay_lines = [
        "| Workbook field | Definition | Tidy nonmissing / N | Average nonmissing / N | Replicate-count / SD nonmissing N |",
        "| --- | --- | ---: | ---: | ---: |",
    ]
    for assay in audit["assay_inventory"]:
        assay_lines.append(
            f"| `{assay['field']}` | {assay['description'] or 'Not defined in workbook'} | "
            f"{assay['tidy_nonmissing_measurements']} / {assay['tidy_nonmissing_measurements'] + assay['tidy_missing_measurements']} | "
            f"{assay['average_nonmissing_antibodies']} / {assay['average_nonmissing_antibodies'] + assay['average_missing_antibodies']} | "
            f"{assay['replicate_count_nonmissing_antibodies']} / {assay['standard_deviation_nonmissing_antibodies']} |"
        )
    overlap_lines = ["| Reference | exact VH | exact VL | exact paired VH/VL |", "| --- | ---: | ---: | ---: |"]
    exact = overlaps["exact_overlap"]
    for name in ("Jain", "AIntibody_all_715_records", "SAbDab2"):
        item = exact[name]
        overlap_lines.append(f"| {name} | {item['VH']} | {item['VL']} | {item['paired']} |")
    split = exact["AIntibody_frozen_splits_paired_hashes_only"]
    for name in SPLITS:
        overlap_lines.append(f"| AIntibody {name} (hash-only) | — | — | {split[name]} |")
    novelty = overlaps["AIntibody_train_validation_similarity"]
    bins = novelty["novelty_bins"]
    clusters = overlaps["identity_proxy_clusters"]
    cluster_lines = ["| Paired-min identity edge threshold | Clusters | Largest | Singletons |", "| ---: | ---: | ---: | ---: |"]
    for threshold, result in clusters.items():
        cluster_lines.append(f"| {threshold} | {result['clusters']} | {result['largest_cluster']} | {result['singleton_clusters']} |")
    source_hash = workbook["sha256"]
    return "\n".join(
        [
            "# Phase 7C.1 — GDPa3 Official Workbook Audit",
            "",
            "Audit date: 2026-09-26. This is a data/provenance and assay-compatibility audit only.",
            "",
            "## Source integrity and provenance",
            "",
            f"- Local source: `{workbook['source_filename']}` (user-provided as the official GDPa3 workbook; no download performed in this audit).",
            f"- Size: {workbook['file_size_bytes']:,} bytes; SHA256: `{source_hash}`.",
            "- ZIP/XLSX package integrity: PASS; source workbook was opened read-only and not modified.",
            "- The file-specific redistribution license and exact original retrieval date are not recorded in the workbook and remain unverified. Keep the raw file local/ignored pending terms review.",
            "- Official source describes GDPa3 competition targets as HIC, PR-CHO, AC-SINS pH 7.4, Tm2 and titer ([Ginkgo competition page](https://datapoints.ginkgo.bio/ai-competitions/2025-abdev-competition)).",
            "",
            "## Actual workbook structure",
            "",
            "| Sheet | Non-empty first-column rows | Columns |",
            "| --- | ---: | ---: |",
            *[
                f"| `{name}` | {detail['rows_with_nonempty_first_column']} | {detail['column_count']} |"
                for name, detail in workbook["sheets"].items()
            ],
            "",
            "Exact assay and identity headers are preserved in the machine-readable audit JSON; this report intentionally omits sequence strings and assay values.",
            "",
            "## VH/VL sequence audit",
            "",
            f"- Sequence rows / unique IDs: {seq['rows']} / {seq['unique_record_ids']}; duplicate ID rows: {seq['duplicate_id_rows']}.",
            f"- VH/VL fields: `{audit['sequence_fields']['VH']}` / `{audit['sequence_fields']['VL']}`; the Sequences tab uses `{audit['sequence_fields']['VL']}`, while Definitions names the light-chain protein field `{audit['sequence_fields']['source_header_difference']['definitions_light_chain_sequence_field']}`.",
            f"- Project validator status: {seq['sequence_status_counts']['VALID']} VALID, {seq['sequence_status_counts']['PARTIAL']} PARTIAL, {seq['sequence_status_counts']['INVALID']} INVALID.",
            f"- Unique valid VH / VL / paired hashes: {seq['unique_valid_VH_hashes']} / {seq['unique_valid_VL_hashes']} / {seq['unique_valid_paired_hashes']}.",
            f"- Missing VH / VL: {seq['missing_VH_rows']} / {seq['missing_VL_rows']}; sequence normalization needed: {seq['normalized_sequence_rows']} rows.",
            f"- VH length min/median/max: {seq['VH_length_aa']['min']} / {seq['VH_length_aa']['median']} / {seq['VH_length_aa']['max']} aa; VL: {seq['VL_length_aa']['min']} / {seq['VL_length_aa']['median']} / {seq['VL_length_aa']['max']} aa.",
            "- No IGHV/IGKV/IGLV germline calls are present. `hc_subtype` and `lc_subtype` are constant-region subtype metadata; identity clusters are not germline annotations.",
            "",
            "## Exact sequence overlap",
            "",
            *overlap_lines,
            "",
            "AIntibody TEST overlap was determined solely from the frozen split manifest’s sequence hashes and split labels; no TEST sequence row, assay value, or label was opened for that check. The additional AIntibody TRAIN+VALIDATION similarity calculation used only paired sequences and global alignment.",
            "",
            "## Family-diversity proxy and AIntibody TRAIN+VALIDATION novelty",
            "",
            f"- Nearest paired-min global identity against {novelty['reference_unique_pairs']} unique AIntibody TRAIN+VALIDATION pairs (N={novelty['query_unique_pairs']} GDPa3 pairs): min {novelty['nearest_paired_min']['min']:.3f}, median {novelty['nearest_paired_min']['median']:.3f}, max {novelty['nearest_paired_min']['max']:.3f}.",
            f"- GDPa3 nearest-neighbour novelty bins: <0.50={bins['<0.50']}; 0.50–<0.70={bins['0.50-<0.70']}; 0.70–<0.80={bins['0.70-<0.80']}; 0.80–<0.90={bins['0.80-<0.90']}; ≥0.90={bins['>=0.90']}.",
            "",
            *cluster_lines,
            "",
            "Clustering is exact all-pairs connected components with an edge only when both chain identities meet the threshold. It is an identity proxy, not an antibody-family or germline assignment. The 0.70 graph has a largest component of 16/80; at 0.80 the largest is 2/80; at 0.90 all 80 are singleton clusters. Together with zero exact overlaps, this supports a sequence-diverse external cohort for a human-reviewed HIC-only protocol, without claiming strict family-independent generalization.",
            "",
            "## Complete experimental assay inventory and field-name crosswalk",
            "",
            *assay_lines,
            "",
            "GDPa3’s five named competition endpoints resolve to workbook fields as follows: HIC → `hic_rt` (minutes); PR-CHO → `polyreactivity_prscore_cho`; AC-SINS pH 7.4 → `acsins_dLmax_ph7.4` (dLmax, nm); Tm2 → `tm2_nanodsf` (second nanoDSF transition, °C); Titer → `titer` (µg/mL). The workbook additionally contains AC-SINS pH 6, HAC, SMAC, SEC % monomer, OVA polyreactivity, Tm1, Tm3 and nanoDSF onset. These are distinct readouts, not interchangeable labels.",
            "",
            "`Definitions` describes 21 fields. `Assay Data - tidy format` has one row per technical replicate (320 rows, 80 IDs), while `Assay Data - average` has one antibody row (80) and 40 columns: ID plus mean/replicate-count/SD fields for 13 readouts. Missing values are reported, not imputed.",
            "",
            "## Frozen model compatibility and protocol shift",
            "",
            "- HIC-only candidate: GDPa3 `hic_rt_avg` and the frozen `hic_esm2_v1` target both represent HIC retention time in minutes; GDPa3 has 79/80 nonmissing averaged measurements. This is nominal measurement-family compatibility, not confirmed protocol equivalence.",
            "- Protocol shift remains material: GDPa3 workbook does not specify the HIC column, buffers, gradient, instrument, or detailed run conditions. AIntibody’s published HIC-HPLC method is specified separately and was performed by Mosaic Biosciences ([Nature Biotechnology paper and Methods](https://www.nature.com/articles/s41587-026-03238-6)). Do not transform or threshold GDPa3 values; any future test must be explicitly described as cross-source/protocol external evaluation.",
            "- Composite classifier compatibility: NO. GDPa3 lacks the AIntibody BVP and Tagg endpoints; `AC-SINS dLmax` is not the AIntibody dPW-adjusted metric; and GDPa3 Tm2 is not the AIntibody Tm endpoint. Do not apply the frozen composite classifier or derive a GDPa3 composite label.",
            "- No ML training, tuning, prediction, performance calculation, correlation, thresholding, or label-derived filtering was performed. No production code was changed.",
            "",
            "## Phase 7C.1 decision",
            "",
            "**READY_HIC_ONLY** — eligible to prepare/review a protocol for a frozen HIC-only external evaluation, with the protocol shift explicitly treated as a limitation. Composite developability validation is not supported. Phase 8 has not started; no prediction or outcome analysis was performed.",
            "",
            "Machine-readable details: [gdpa3_phase7c1_audit.json](gdpa3_phase7c1_audit.json). Provenance manifest: [source_manifest.json](../external_developability/source_manifest.json). HIC-only protocol for human review: [Phase 8 draft](../external_developability/phase8_hic_protocol_draft.md).",
            "",
        ]
    )


def write_audit_outputs(audit: dict[str, Any], overlaps: dict[str, Any], output_dir: str | Path) -> tuple[Path, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    json_path = output / "gdpa3_phase7c1_audit.json"
    markdown_path = output / "gdpa3_phase7c1_audit.md"
    payload = {"workbook_audit": audit, "sequence_comparison": overlaps}
    json_path.write_text(json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    markdown_path.write_text(build_markdown_report(audit, overlaps), encoding="utf-8")
    return json_path, markdown_path

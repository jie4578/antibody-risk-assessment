"""Inspect the official AIntibody Supplementary Data 1–3 workbook."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl


def inspect_aintibody(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 4), values_only=True))
        sheets.append({
            "name": sheet.title,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "first_rows": [[None if value is None else str(value) for value in row] for row in rows],
        })
    return {"dataset": "aintibody_2026", "file": str(path).replace("\\", "/"), "sheets": sheets}


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    path = Path(__file__).resolve().parents[1] / "data" / "raw" / "aintibody_2026" / "41587_2026_3238_MOESM4_ESM.xlsx"
    print(json.dumps(inspect_aintibody(path), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

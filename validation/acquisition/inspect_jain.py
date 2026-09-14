"""Inspect Jain et al. 2017 official supplementary workbooks.

The inspector reports structure only. It does not assign scientific meaning
to an assay column until the source workbook is available and reviewed.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl


FILES = {
    "S1": "pnas.1616408114.sd01.xlsx",
    "S2": "pnas.1616408114.sd02.xlsx",
    "S3": "pnas.1616408114.sd03.xlsx",
}


def inspect_workbook(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"file": str(path).replace("\\", "/"), "status": "missing"}
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return {"file": str(path).replace("\\", "/"), "status": "unreadable", "error": str(exc)}
    sheets = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8), values_only=True))
        sheets.append({
            "name": sheet.title,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "first_rows": [[None if value is None else str(value) for value in row] for row in rows],
        })
    return {"file": str(path).replace("\\", "/"), "status": "ok", "sheets": sheets}


def inspect_jain(raw_dir: str | Path) -> dict[str, Any]:
    root = Path(raw_dir)
    return {"dataset": "jain_2017", "files": [inspect_workbook(root / name) for name in FILES.values()]}


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    root = Path(__file__).resolve().parents[1] / "data" / "raw" / "jain_2017"
    print(json.dumps(inspect_jain(root), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
